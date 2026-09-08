"""Stream the public Kalshi census and raw candle extracts into partitioned Excel.

Run with the bundled Python containing openpyxl. Raw gzip artifacts are authoritative;
Excel is a typed, auditable presentation with documented numeric/text limitations.
"""
from __future__ import annotations

import argparse
import gzip
import json
import math
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

UTC = timezone.utc
DATE_FORMAT = 'yyyy-mm-dd hh:mm:ss" UTC"'
NUMBER_FORMAT = '#,##0.########'
MONEY_FORMAT = '"$"#,##0.0000'
BAD_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
SOURCE_URL = 'https://api.elections.kalshi.com/trade-api/v2'
MAX_TEXT = 32767


def utc_now() -> str:
    return datetime.now(UTC).isoformat().replace('+00:00', 'Z')


def read_json(path: Path) -> Any:
    opener = gzip.open if path.suffix == '.gz' else open
    with opener(path, 'rt', encoding='utf-8') as stream:
        return json.load(stream)


def decimal(value: Any) -> Decimal | None:
    if value is None or value == '' or isinstance(value, bool):
        return None
    try:
        result = Decimal(str(value))
        return result if result.is_finite() else None
    except (ValueError, InvalidOperation):
        return None


def preferred(raw: dict, *keys: str) -> Any:
    for key in keys:
        if key in raw and raw[key] is not None:
            return raw[key]
    return None


def units(raw: dict, dollars: str, cents: str) -> Decimal | None:
    if dollars in raw and raw[dollars] is not None:
        return decimal(raw[dollars])
    value = decimal(raw.get(cents))
    return value / 100 if value is not None else None


def product(a: Any, b: Any) -> Decimal | None:
    a, b = decimal(a), decimal(b)
    return a * b if a is not None and b is not None else None


def as_datetime(value: Any) -> datetime | None:
    if value is None or value == '':
        return None
    try:
        if isinstance(value, (int, float, Decimal)):
            result = datetime.fromtimestamp(float(value), UTC)
        else:
            result = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
            if result.tzinfo is None:
                return None
        return result.astimezone(UTC).replace(tzinfo=None)
    except (ValueError, TypeError, OverflowError, OSError):
        return None


def json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'), default=str)


def literal(value: str) -> tuple[str, list[str]]:
    flags = []
    clean = BAD_XML.sub('', value)
    if clean != value:
        flags.append('xml_controls_removed_raw_preserved')
    if len(clean) > MAX_TEXT:
        clean = clean[:MAX_TEXT - 45] + ' [TRUNCATED; complete text in raw artifact]'
        flags.append('excel_text_truncated_raw_preserved')
    # Force string cells in writer; no user-controlled text becomes a formula.
    return clean, flags


def time_field(key: str) -> bool:
    return key.endswith(('_time', '_at', '_utc')) or key in {'asof_utc', 'end_period_utc', 'start_period_utc'}


def numeric_field(key: str) -> bool:
    return key.endswith(('_fp', '_contracts', '_usd', '_dollars')) or key in {
        'volume', 'volume_24h', 'open_interest', 'liquidity', 'notional_value',
        'yes_bid', 'yes_ask', 'no_bid', 'no_ask', 'last_price', 'previous_price',
        'previous_yes_bid', 'previous_yes_ask', 'floor_strike', 'cap_strike',
        'yes_bid_size', 'yes_ask_size', 'no_bid_size', 'no_ask_size',
        'tick_size', 'expiration_value', 'settlement_value', 'response_price_units',
    }


def value_for_excel(value: Any, key: str) -> tuple[Any, list[str]]:
    flags = []
    if isinstance(value, (dict, list, tuple)):
        return literal(json_text(value))
    if isinstance(value, datetime):
        return value.replace(tzinfo=None), flags
    if isinstance(value, str) and time_field(key):
        parsed = as_datetime(value)
        if parsed is not None:
            return parsed, flags
    if isinstance(value, str) and numeric_field(key):
        parsed = decimal(value)
        if parsed is not None:
            value = parsed
    if isinstance(value, Decimal):
        if len(value.as_tuple().digits) > 15:
            flags.append('excel_numeric_precision_raw_preserved')
        value = float(value)
    if isinstance(value, float) and not math.isfinite(value):
        return None, ['nonfinite_number_raw_preserved']
    if isinstance(value, int) and not isinstance(value, bool) and abs(value) >= 10**15:
        return str(value), ['large_integer_stored_as_text']
    if isinstance(value, str):
        return literal(value)
    return value, flags


MASTER_PREFIX = [
    'market_ticker', 'series_ticker', 'event_ticker', 'category', 'category_paths',
    'membership_confidence', 'scope_flags', 'eligibility_class', 'eligibility_reason', 'source_tier',
    'collection_status', 'snapshot_observed_utc', 'snapshot_time_basis', 'face_value_usd',
    'volume_contracts', 'volume_24h_contracts', 'open_interest_contracts',
    'volume_face_notional_usd', 'volume_24h_face_notional_usd',
    'open_interest_face_notional_usd', 'yes_bid_usd', 'yes_ask_usd',
    'no_bid_usd', 'no_ask_usd', 'last_price_usd', 'quoted_midpoint_usd',
    'quoted_spread_usd', 'quote_quality', 'raw_artifact',
]
CANDLE_COLUMNS = [
    'market_ticker', 'series_ticker', 'category', 'window_label', 'source_tier',
    'collection_status', 'period_interval_minutes', 'end_period_ts',
    'start_period_utc', 'end_period_utc', 'partial_start', 'partial_end',
    'boundary_only', 'window_overlap_start_utc', 'window_overlap_end_utc',
    'yes_bid_open_usd', 'yes_bid_high_usd', 'yes_bid_low_usd', 'yes_bid_close_usd',
    'yes_ask_open_usd', 'yes_ask_high_usd', 'yes_ask_low_usd', 'yes_ask_close_usd',
    'price_open_usd', 'price_high_usd', 'price_low_usd', 'price_close_usd',
    'price_mean_usd', 'price_previous_usd', 'volume_contracts',
    'open_interest_end_contracts', 'face_value_usd', 'volume_face_notional_usd',
    'open_interest_end_face_notional_usd', 'retrieved_at', 'data_flags', 'raw_artifact',
]
WINDOW_COLUMNS = [
    'market_ticker', 'series_ticker', 'dataset', 'label', 'period_interval',
    'status', 'rows', 'requested_start_ts', 'requested_end_ts',
    'effective_start_ts', 'effective_end_ts', 'first_end_period_ts',
    'last_end_period_ts', 'errors', 'requests', 'extra_fields', 'raw_artifact',
]
UTC_DAILY_COLUMNS = [
    'market_ticker', 'series_ticker', 'category', 'window_label', 'source_tier',
    'collection_status', 'utc_day', 'start_period_utc', 'end_period_utc',
    'nominal_hours', 'expected_intersecting_hours', 'expected_full_hours',
    'observed_hours', 'missing_intersecting_hours', 'volume_observation_hours',
    'missing_price_hours', 'duplicate_hour_timestamps', 'observed_volume_contracts',
    'open_interest_latest_contracts', 'open_interest_latest_utc', 'oi_at_day_end',
    'face_value_usd', 'observed_volume_face_notional_usd',
    'open_interest_latest_face_notional_usd', 'price_open_usd', 'price_high_usd',
    'price_low_usd', 'price_close_usd', 'partial_start', 'partial_end',
    'volume_complete', 'data_flags', 'raw_artifact',
]
BOOK_LEVEL_COLUMNS = [
    'market_ticker', 'series_ticker', 'category', 'retrieved_at', 'snapshot_time_basis',
    'collection_status', 'raw_side', 'raw_level_index', 'bid_price_usd',
    'quantity_contracts', 'derived_opposite_ask_side', 'derived_opposite_ask_usd',
    'face_value_usd', 'data_flags', 'raw_artifact', 'raw_response_path',
]
BOOK_SUMMARY_COLUMNS = [
    'market_ticker', 'series_ticker', 'category', 'retrieved_at', 'snapshot_time_basis',
    'collection_status', 'face_value_usd', 'yes_bid_usd', 'yes_ask_usd',
    'yes_bid_top_contracts', 'yes_ask_top_contracts', 'quoted_spread_usd',
    'quoted_midpoint_usd', 'yes_bid_observed_depth_contracts',
    'yes_ask_observed_depth_contracts', 'yes_bid_levels', 'yes_ask_levels',
    'data_flags', 'raw_artifact', 'raw_response_path',
]


class PartitionWriter:
    """Bound memory and Excel row limits; save closed partitions immediately."""
    def __init__(self, out: Path, dataset: str, columns: list[str], state: str,
                 asof: str, max_sheet: int = 350000, max_book: int = 500000):
        self.out, self.dataset, self.columns = out, dataset, columns
        self.state, self.asof = state, asof
        self.max_sheet, self.max_book = max_sheet, max_book
        self.wb = self.ws = None
        self.sheet_rows = self.book_rows = self.total = self.part = self.sheet_n = 0
        self.files: list[dict] = []
        self.flags: Counter = Counter()

    def _create(self):
        self.part += 1
        self.book_rows = self.sheet_n = 0
        self.wb = Workbook(write_only=True)
        self.wb.properties.title = f'Kalshi {self.dataset}'
        self.wb.properties.subject = self.state
        self.wb.properties.creator = 'Public Kalshi API extraction'
        self.wb.properties.description = 'Source data and exact precision retained in gzip JSON artifacts.'
        self._sheet()

    def _sheet(self):
        self.sheet_n += 1
        self.sheet_rows = 0
        self.ws = self.wb.create_sheet(f'Data_{self.sheet_n:02d}')
        self.ws.freeze_panes = 'C4'
        self.ws.sheet_view.showGridLines = False
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.ws.print_options.horizontalCentered = False
        self.ws.print_title_rows = '1:3'
        heads = []
        for i, name in enumerate(self.columns + ['excel_flags'], 1):
            cell = WriteOnlyCell(self.ws, name)
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='203864')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            heads.append(cell)
            width = 23 if ('utc' in name or time_field(name)) else min(36, max(16, len(name) * .9))
            if name in {'rules_primary', 'rules_secondary', 'title', 'subtitle', 'raw_artifact'}:
                width = 60
            self.ws.column_dimensions[get_column_letter(i)].width = width
        self.ws.row_dimensions[3].height = 32
        self.ws.append([f'Kalshi {self.dataset}', self.state])
        self.ws.append(['As of UTC', self.asof, 'Nulls are blank. UTC times. Exact raw data retained. See INDEX.xlsx.'])
        self.ws.append(heads)

    def _end_sheet(self):
        if self.ws is not None:
            self.ws.auto_filter.ref = f'A3:{get_column_letter(len(self.columns)+1)}{self.sheet_rows+3}'

    def _save(self):
        if self.wb is None:
            return
        self._end_sheet()
        path = self.out / f'{self.dataset}_{self.part:04d}.xlsx'
        pending = path.with_suffix('.xlsx.tmp')
        self.wb.save(pending)
        pending.replace(path)
        self.files.append({'dataset': self.dataset, 'file': path.name,
                           'rows': self.book_rows, 'sheets': self.sheet_n,
                           'bytes': path.stat().st_size, 'status': self.state})
        print(json_text({'saved': str(path), 'rows': self.book_rows}), flush=True)
        self.wb = self.ws = None

    def append(self, row: dict):
        if self.wb is None:
            self._create()
        if self.book_rows >= self.max_book:
            self._save()
            self._create()
        elif self.sheet_rows >= self.max_sheet:
            self._end_sheet()
            self._sheet()
        cells, flags = [], []
        for key in self.columns:
            value, additional = value_for_excel(row.get(key), key)
            flags.extend(additional)
            if value is None:
                cells.append(None)
                continue
            cell = WriteOnlyCell(self.ws, value=value)
            if isinstance(value, str):
                cell.data_type = 's'
            elif isinstance(value, datetime):
                cell.number_format = DATE_FORMAT
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = MONEY_FORMAT if key.endswith(('_usd', '_dollars')) else NUMBER_FORMAT
            cells.append(cell)
        flags = sorted(set(flags))
        self.flags.update(flags)
        flag_cell = WriteOnlyCell(self.ws, '; '.join(flags))
        flag_cell.data_type = 's'
        cells.append(flag_cell)
        self.ws.append(cells)
        self.sheet_rows += 1
        self.book_rows += 1
        self.total += 1

    def close(self) -> list[dict]:
        self._save()
        return self.files


def master_row(market: dict, payload: dict, source: str) -> dict:
    series = payload.get('series') or {}
    face = units(market, 'notional_value_dollars', 'notional_value')
    vol = decimal(preferred(market, 'volume_fp', 'volume'))
    dayvol = decimal(preferred(market, 'volume_24h_fp', 'volume_24h'))
    oi = decimal(preferred(market, 'open_interest_fp', 'open_interest'))
    bid = units(market, 'yes_bid_dollars', 'yes_bid')
    ask = units(market, 'yes_ask_dollars', 'yes_ask')
    empty = any(decimal(market.get(k)) == 0 for k in ('yes_bid_size_fp', 'yes_bid_size', 'yes_ask_size_fp', 'yes_ask_size') if market.get(k) is not None)
    valid = bid is not None and ask is not None and face is not None and Decimal(0) <= bid <= ask <= face and not empty
    quality = 'quoted_two_sided_depth_not_verified' if valid else 'missing_crossed_out_of_range_or_empty_quote'
    row = dict(market)
    row.update({
        'market_ticker': market.get('ticker'), 'series_ticker': payload.get('series_ticker'),
        'category': series.get('category'), 'category_paths': payload.get('category_paths'),
        'membership_confidence': payload.get('membership_confidence'),
        'scope_flags': payload.get('scope_flags') or series.get('_scope_flags'),
        'eligibility_class': market.get('_eligibility_class'),
        'eligibility_reason': market.get('_eligibility_reason'),
        'source_tier': market.get('_source_tier'), 'collection_status': payload.get('status'),
        'snapshot_observed_utc': market.get('_retrieved_at'),
        'snapshot_time_basis': 'market_raw_page_retrieval' if market.get('_retrieved_at') else 'unavailable_pending_raw_page_enrichment', 'face_value_usd': face,
        'volume_contracts': vol, 'volume_24h_contracts': dayvol, 'open_interest_contracts': oi,
        'volume_face_notional_usd': product(vol, face),
        'volume_24h_face_notional_usd': product(dayvol, face),
        'open_interest_face_notional_usd': product(oi, face),
        'yes_bid_usd': bid, 'yes_ask_usd': ask,
        'no_bid_usd': units(market, 'no_bid_dollars', 'no_bid'),
        'no_ask_usd': units(market, 'no_ask_dollars', 'no_ask'),
        'last_price_usd': units(market, 'last_price_dollars', 'last_price'),
        'quoted_midpoint_usd': (bid + ask) / 2 if valid else None,
        'quoted_spread_usd': ask - bid if valid else None, 'quote_quality': quality,
        'raw_artifact': source,
    })
    return row


def candle_price(bucket: dict, field: str, source_tier: str, explicit_unit: str | None) -> Decimal | None:
    if f'{field}_dollars' in bucket:
        return decimal(bucket.get(f'{field}_dollars'))
    value = decimal(bucket.get(field))
    if value is None:
        return None
    # Historical endpoint's plain OHLC fields are dollar strings. Legacy live
    # endpoint's integer OHLC fields are cents. Units follow source schema.
    return value if explicit_unit == 'dollars' or source_tier == 'historical' else value / 100


def candle_row(candle: dict, payload: dict, interval: int, source: str,
               metadata: dict | None, default_window: str) -> dict:
    metadata = metadata or {}
    tier = candle.get('_source_tier') or payload.get('source_tier')
    end = candle.get('end_period_ts')
    end_dec = decimal(end)
    face = decimal(payload.get('notional_value_dollars'))
    if face is None:
        face = decimal(metadata.get('face_value_usd'))
    volume = decimal(preferred(candle, 'volume_fp', 'volume'))
    oi = decimal(preferred(candle, 'open_interest_fp', 'open_interest'))
    flags = []
    if face is None:
        flags.append('face_value_unavailable')
    row = {
        'market_ticker': payload.get('market_ticker'), 'series_ticker': payload.get('series_ticker'),
        'category': metadata.get('category'), 'window_label': candle.get('_window_label') or default_window,
        'source_tier': tier, 'collection_status': payload.get('status'),
        'period_interval_minutes': interval, 'end_period_ts': end,
        'start_period_utc': as_datetime(candle.get('_period_start_ts')) if candle.get('_period_start_ts') is not None else as_datetime(end_dec - interval * 60) if end_dec is not None else None,
        'end_period_utc': as_datetime(end), 'partial_start': candle.get('_partial_start'),
        'partial_end': candle.get('_partial_end'), 'boundary_only': candle.get('_boundary_only'),
        'window_overlap_start_utc': as_datetime(candle.get('_window_overlap_start_ts')),
        'window_overlap_end_utc': as_datetime(candle.get('_window_overlap_end_ts')), 'volume_contracts': volume,
        'open_interest_end_contracts': oi, 'face_value_usd': face,
        'volume_face_notional_usd': product(volume, face),
        'open_interest_end_face_notional_usd': product(oi, face),
        'retrieved_at': candle.get('_retrieved_at') or payload.get('retrieved_at'),
        'raw_artifact': source,
    }
    for bucket_name in ('yes_bid', 'yes_ask', 'price'):
        bucket = candle.get(bucket_name) or {}
        for field in ('open', 'high', 'low', 'close', 'mean', 'previous'):
            key = f'{bucket_name}_{field}_usd'
            if key in CANDLE_COLUMNS:
                row[key] = candle_price(bucket, field, tier, candle.get('_price_unit'))
        if bucket_name == 'price' and all(row.get(f'price_{field}_usd') is None for field in ('open', 'high', 'low', 'close')):
            flags.append('no_trade_price_in_candle')
    if face is not None and any(v is not None and (v < 0 or v > face) for k, v in row.items() if k.startswith(('price_', 'yes_bid_', 'yes_ask_')) and k.endswith('_usd')):
        flags.append('price_outside_face_bounds_raw_preserved')
    row['data_flags'] = '; '.join(flags)
    return row


def explicit_complete(manifest: dict | None) -> bool:
    if not manifest or manifest.get('errors'):
        return False
    return manifest.get('complete') is True or manifest.get('status') == 'complete'


def utc_daily_rows(payload: dict, source: str, metadata: dict | None,
                   default_window: str) -> Iterable[dict]:
    """Aggregate observed hourly data; never assume sparse intervals are zero."""
    groups = defaultdict(list)
    for raw in payload.get('hourly') or []:
        end = decimal(raw.get('end_period_ts'))
        if end is not None:
            day_start = int((end - 1) // 86400) * 86400
            groups[(raw.get('_window_label') or default_window, day_start)].append(raw)
    for (label, start), raw_rows in sorted(groups.items()):
        window = next((w for w in payload.get('windows') or [] if w.get('label') == label and w.get('period_interval') == 60), {})
        effective_start = decimal(window.get('effective_start_ts'))
        effective_end = decimal(window.get('effective_end_ts'))
        lo = max(Decimal(start), effective_start) if effective_start is not None else None
        hi = min(Decimal(start + 86400), effective_end) if effective_end is not None else None
        expected = max(0, math.ceil(hi / 3600) - math.floor(lo / 3600)) if lo is not None and hi is not None and hi > lo else None
        full_hours = max(0, math.floor(hi / 3600) - math.ceil(lo / 3600)) if lo is not None and hi is not None and hi > lo else None
        by_end = {}
        duplicates = 0
        for raw in raw_rows:
            end = raw['end_period_ts']
            if end in by_end:
                duplicates += 1
            else:
                by_end[end] = raw
        rows = [candle_row(raw, payload, 60, source, metadata, label) for _, raw in sorted(by_end.items(), key=lambda item: float(item[0]))]
        volumes = [r['volume_contracts'] for r in rows if r['volume_contracts'] is not None]
        observed_volume = sum(volumes, Decimal(0)) if volumes else None
        last_oi = next((r for r in reversed(rows) if r['open_interest_end_contracts'] is not None), None)
        partial_start = any(bool(r['partial_start']) for r in rows)
        partial_end = any(bool(r['partial_end']) for r in rows)
        observed = len(rows)
        missing = max(0, expected - observed) if expected is not None else None
        flags = []
        if missing:
            flags.append('sparse_or_unfinished_hours_not_filled')
        if expected is None:
            flags.append('requested_window_unavailable')
        if partial_start or partial_end:
            flags.append('full_boundary_candle_not_prorated')
        if duplicates:
            flags.append('duplicate_hours_first_retained')
        face = rows[0]['face_value_usd']
        oi = last_oi['open_interest_end_contracts'] if last_oi else None
        row = {
            'market_ticker': payload.get('market_ticker'), 'series_ticker': payload.get('series_ticker'),
            'category': (metadata or {}).get('category'), 'window_label': label,
            'source_tier': payload.get('source_tier'), 'collection_status': payload.get('status'),
            'utc_day': as_datetime(start), 'start_period_utc': as_datetime(start),
            'end_period_utc': as_datetime(start + 86400), 'nominal_hours': 24,
            'expected_intersecting_hours': expected, 'expected_full_hours': full_hours,
            'observed_hours': observed, 'missing_intersecting_hours': missing,
            'volume_observation_hours': len(volumes),
            'missing_price_hours': sum(all(r.get(f'price_{f}_usd') is None for f in ('open', 'high', 'low', 'close')) for r in rows),
            'duplicate_hour_timestamps': duplicates, 'observed_volume_contracts': observed_volume,
            'open_interest_latest_contracts': oi,
            'open_interest_latest_utc': last_oi['end_period_utc'] if last_oi else None,
            'oi_at_day_end': last_oi['end_period_utc'] == as_datetime(start + 86400) if last_oi else False,
            'face_value_usd': face, 'observed_volume_face_notional_usd': product(observed_volume, face),
            'open_interest_latest_face_notional_usd': product(oi, face),
            'partial_start': partial_start, 'partial_end': partial_end,
            'volume_complete': expected is not None and missing == 0 and len(volumes) == observed and not duplicates and not partial_start and not partial_end,
            'data_flags': '; '.join(flags), 'raw_artifact': source,
        }
        for field in ('open', 'high', 'low', 'close'):
            values = [r[f'price_{field}_usd'] for r in rows if r.get(f'price_{field}_usd') is not None]
            row[f'price_{field}_usd'] = (values[0] if field == 'open' else values[-1] if field == 'close' else max(values) if field == 'high' else min(values)) if values else None
        yield row


def book_rows(payload: dict, source: str, metadata: dict | None) -> tuple[dict, list[dict]]:
    metadata = metadata or {}
    face = decimal(metadata.get('face_value_usd'))
    container = payload.get('orderbook') or {}
    book = container.get('orderbook_fp') or container.get('orderbook') or {}
    context = {
        'market_ticker': payload.get('ticker'), 'series_ticker': payload.get('series_ticker'),
        'category': metadata.get('category'), 'retrieved_at': payload.get('retrieved_at'),
        'snapshot_time_basis': payload.get('snapshot_time_basis'),
        'collection_status': payload.get('status'), 'face_value_usd': face,
        'raw_artifact': source, 'raw_response_path': payload.get('raw_response_path'),
    }
    levels, valid = [], {'yes': [], 'no': []}
    flags = set()
    for side in ('yes', 'no'):
        source_levels = book.get(f'{side}_dollars') or []
        for i, item in enumerate(source_levels, 1):
            rowflags = []
            price = decimal(item[0]) if isinstance(item, list) and len(item) >= 2 else None
            qty = decimal(item[1]) if isinstance(item, list) and len(item) >= 2 else None
            if price is None or qty is None or qty <= 0 or face is None or not Decimal(0) <= price <= face:
                rowflags.append('invalid_or_unverifiable_level_excluded_from_summary')
            else:
                valid[side].append((price, qty))
            flags.update(rowflags)
            levels.append({**context, 'raw_side': f'{side}_bid', 'raw_level_index': i,
                           'bid_price_usd': price, 'quantity_contracts': qty,
                           'derived_opposite_ask_side': 'no_ask' if side == 'yes' else 'yes_ask',
                           'derived_opposite_ask_usd': face - price if face is not None and price is not None else None,
                           'data_flags': '; '.join(rowflags)})
    ybest = max((p for p, q in valid['yes']), default=None)
    nbest = max((p for p, q in valid['no']), default=None)
    ask = face - nbest if face is not None and nbest is not None else None
    paired = ybest is not None and ask is not None and ybest <= ask
    if not paired:
        flags.add('one_sided_empty_or_crossed_book')
    summary = {**context, 'yes_bid_usd': ybest, 'yes_ask_usd': ask,
               'yes_bid_top_contracts': sum((q for p, q in valid['yes'] if p == ybest), Decimal(0)) if ybest is not None else None,
               'yes_ask_top_contracts': sum((q for p, q in valid['no'] if p == nbest), Decimal(0)) if nbest is not None else None,
               'quoted_spread_usd': ask - ybest if paired else None,
               'quoted_midpoint_usd': (ask + ybest) / 2 if paired else None,
               'yes_bid_observed_depth_contracts': sum((q for p, q in valid['yes']), Decimal(0)) if valid['yes'] else None,
               'yes_ask_observed_depth_contracts': sum((q for p, q in valid['no']), Decimal(0)) if valid['no'] else None,
               'yes_bid_levels': len(valid['yes']), 'yes_ask_levels': len(valid['no']),
               'data_flags': '; '.join(sorted(flags))}
    return summary, levels


def assess_state(root: Path, errors: list[dict], inputs: dict) -> tuple[str, dict]:
    docs = {}
    for name in ('run_config.json', 'discovery_manifest.json', 'candle_manifest.json', 'candles_manifest.json',
                 'collection_manifest.json', 'census_progress.json', 'candles_progress.json', 'books_manifest.json'):
        path = root / name
        if path.exists():
            try:
                docs[name] = read_json(path)
            except Exception as exc:
                errors.append({'artifact': name, 'error': str(exc)})
    final = docs.get('collection_manifest.json') or {}
    if not errors and explicit_complete(final) and all(final.get(k) is True for k in ('census_complete', 'category_membership_complete', 'candles_complete')) and inputs.get('datasets_complete'):
        return 'FULL WITHIN DOCUMENTED SCOPE', docs
    if not errors and explicit_complete(final) and final.get('census_complete') is True and final.get('candles_complete') is True and inputs.get('datasets_complete'):
        return 'API EXTRACT COMPLETE; CATEGORY SCOPE APPROXIMATE', docs
    return 'PARTIAL EXTRACT — SEE COVERAGE', docs


def dictionary_rows(master_columns: list[str]) -> list[dict]:
    definitions = {
        'face_value_usd': 'Contract face/notional value from API. No assumed $1 when missing.',
        'volume_contracts': 'Contracts traded in candle interval, or lifetime total in market master. Fractional counts supported.',
        'volume_24h_contracts': 'Market snapshot trailing 24-hour traded contract count. Not a calendar-day total.',
        'open_interest_contracts': 'Snapshot outstanding contracts, a stock. Not net bullish positioning or number of traders.',
        'open_interest_end_contracts': 'Outstanding contracts at candle end. A stock: never sum across hours or days.',
        'volume_face_notional_usd': 'Volume contracts × source face value. Face turnover, not premium cash turnover.',
        'volume_24h_face_notional_usd': 'Trailing 24-hour contracts × source face value. Not premium traded.',
        'open_interest_face_notional_usd': 'Snapshot open contracts × source face value. Not directional flow or collateral.',
        'open_interest_end_face_notional_usd': 'Candle ending open interest × source face value. Not summed through time.',
        'quoted_midpoint_usd': 'Mean of valid source Yes bid/ask, with valid face bounds and no explicit zero quote size. Executable depth not verified.',
        'quoted_spread_usd': 'Yes ask less Yes bid under the same quote checks. Dollars per contract.',
        'raw_artifact': 'Path relative to extraction root. The original response carries fields/precision omitted from Excel.',
        'snapshot_observed_utc': 'Retrieval time of corresponding raw market API page. Blank if enrichment is pending; not the frozen census cutoff or last-trade time.',
        'end_period_ts': 'API candle end in Unix seconds. Daily boundaries follow API sessions, not necessarily UTC midnight.',
        'start_period_utc': 'API candle end less requested interval. Boundary candles can overlap requested window; see partial flags.',
        'end_period_utc': 'API candle end converted to a typed UTC date (timezone label in format).',
        'window_label': 'common = requested last-two-month window. presettlement = separate approved extension before settlement.',
        'partial_start': 'Candle begins before effective requested window; full candle retained and not prorated.',
        'partial_end': 'Candle ends after effective requested window; full candle retained and not prorated.',
        'collection_status': 'Source artifact complete/partial/error. Complete retrieval does not guarantee traded observations in every hour.',
        'membership_confidence': 'Discovery classification confidence; API categories may not exactly reproduce website subcategory membership.',
        'eligibility_class': 'Source-enriched class: confirmed_active, recently_settled, unresolved_candidate, or settlement_date_uncertain. Retained unresolved candidates are not automatically trading-active.',
        'excel_flags': 'Text truncation, sanitization or precision caveat affecting this row. Complete original value remains in gzip.',
    }
    rows = []
    definitions.update({
        'utc_day': 'UTC calendar day, grouping each hourly candle by date(end_period_ts minus one second).',
        'observed_volume_contracts': 'Sum of nonmissing observed hourly volumes. Incomplete when hours or volume fields are missing, or boundary candles overlap the requested interval.',
        'open_interest_latest_contracts': 'Latest nonmissing hourly ending OI assigned to this UTC day. A stock; never summed.',
        'open_interest_latest_utc': 'Timestamp of latest observed nonmissing OI used, which may precede day end.',
        'oi_at_day_end': 'True only when OI observation ends exactly at next UTC midnight.',
        'expected_intersecting_hours': 'Number of hourly intervals intersecting effective requested window inside this UTC day, including incomplete boundary hours.',
        'expected_full_hours': 'Number of whole hourly intervals contained in effective requested window inside this UTC day.',
        'observed_hours': 'Distinct hourly candle timestamps actually returned. Missing timestamps are not filled.',
        'volume_complete': 'Observed volume covers expected intersecting hours with nonmissing volumes and no duplicate or straddling boundary candles. Does not imply a full 24-hour calendar day.',
        'raw_side': 'Public order-book bid side. YES and NO are bid ladders; opposite asks are derived from face value.',
        'quantity_contracts': 'Displayed order quantity at this price, not a trader position or identified investor.',
        'snapshot_time_basis': 'Book snapshots use retrieval time, after the fixed census as-of. They are not historical order books.',
    })
    for dataset, columns in [('market_master', master_columns), ('hourly_and_api_daily_sessions', CANDLE_COLUMNS), ('utc_daily', UTC_DAILY_COLUMNS), ('book_levels', BOOK_LEVEL_COLUMNS), ('book_summary', BOOK_SUMMARY_COLUMNS), ('window_coverage', WINDOW_COLUMNS)]:
        for key in columns + ['excel_flags']:
            description = definitions.get(key)
            if not description:
                if key.startswith(('price_', 'yes_bid_', 'yes_ask_')) and key.endswith('_usd'):
                    description = 'Source candle price statistic, normalized to USD by endpoint schema. Null remains missing; no forward fill.'
                elif dataset == 'market_master' and key not in MASTER_PREFIX:
                    description = 'Original top-level market API field; nested structures are JSON text. Field-name units preserved. See raw artifact.'
                else:
                    description = 'Source identity, status, provenance or coverage field. Original response retained.'
            rows.append({'dataset': dataset, 'field': key, 'definition': description})
    return rows


def export(root: Path, out: Path, max_sheet: int, max_book: int, selected: set[str] | None = None) -> dict:
    root, out = root.resolve(), out.resolve()
    out.mkdir(parents=True, exist_ok=True)
    errors: list[dict] = []
    discovery_files = sorted((root / 'discovery').glob('*.json.gz'))
    captured = {folder: sorted((root / folder).glob('*.json.gz')) for folder in ('candles', 'candles_presettlement')}
    captured_books = sorted((root / 'books').glob('*.json.gz'))
    master_keys: set[str] = set()
    metadata: dict[str, dict] = {}
    status_counts: Counter = Counter()
    artifacts: list[dict] = []
    seen = set()
    for path in discovery_files:
        try:
            payload = read_json(path)
        except Exception as exc:
            errors.append({'artifact': str(path.relative_to(root)), 'error': str(exc)})
            continue
        status_counts[f'discovery_{payload.get("status", "unknown")}'] += 1
        if payload.get('status') != 'complete' or payload.get('errors'):
            errors.append({'artifact': str(path.relative_to(root)), 'error': payload.get('errors') or payload.get('status')})
        artifacts.append({'dataset': 'discovery', 'artifact': str(path.relative_to(root)), 'status': payload.get('status'), 'rows': len(payload.get('markets') or []), 'errors': payload.get('errors')})
        for market in payload.get('markets') or []:
            master_keys.update(market)
            ticker = market.get('ticker')
            if ticker in seen:
                errors.append({'artifact': str(path.relative_to(root)), 'error': f'duplicate market ticker {ticker}; first retained'})
                continue
            seen.add(ticker)
            metadata[ticker] = {'category': (payload.get('series') or {}).get('category'), 'face_value_usd': units(market, 'notional_value_dollars', 'notional_value')}
    candle_tickers = set()
    for folder, paths in captured.items():
        for path in paths:
            try:
                payload = read_json(path)
                status_counts[f'{folder}_{payload.get("status", "unknown")}'] += 1
                if folder == 'candles':
                    candle_tickers.add(payload.get('market_ticker'))
                if payload.get('status') != 'complete' or payload.get('errors'):
                    errors.append({'artifact': str(path.relative_to(root)), 'error': payload.get('errors') or payload.get('status')})
                artifacts.append({'dataset': folder, 'artifact': str(path.relative_to(root)), 'status': payload.get('status'), 'rows': len(payload.get('hourly') or []) + len(payload.get('daily') or []), 'errors': payload.get('errors')})
            except Exception as exc:
                errors.append({'artifact': str(path.relative_to(root)), 'error': str(exc)})
    missing = sorted(seen - candle_tickers)
    extra = sorted(candle_tickers - seen)
    inputs = {'discovery_files': len(discovery_files), 'eligible_markets': len(seen), 'candle_files': len(captured['candles']), 'presettlement_files': len(captured['candles_presettlement']), 'missing_candle_markets': missing, 'candle_markets_outside_master': extra, 'status_counts': dict(status_counts), 'datasets_complete': bool(seen) and not missing and not extra and not errors}
    state, docs = assess_state(root, errors, inputs)
    if selected is not None:
        state = 'SELECTED DATASETS — SEE COVERAGE'
    config = docs.get('run_config.json') or {}
    asof = config.get('asof_utc') or 'UNAVAILABLE'
    columns = MASTER_PREFIX + sorted(master_keys - set(MASTER_PREFIX))
    writers = {}

    def writer(name: str, headers: list[str]) -> PartitionWriter:
        if name not in writers:
            writers[name] = PartitionWriter(out, name, headers, state, asof, max_sheet, max_book)
        return writers[name]

    if selected is None or 'market_master' in selected:
        emitted = set()
        for path in discovery_files:
            try:
                payload = read_json(path)
                for market in payload.get('markets') or []:
                    ticker = market.get('ticker')
                    if ticker in emitted:
                        continue
                    emitted.add(ticker)
                    writer('market_master', columns).append(master_row(market, payload, str(path.relative_to(root))))
            except Exception as exc:
                errors.append({'artifact': str(path.relative_to(root)), 'error': str(exc), 'stage': 'master_export'})
        if 'market_master' in writers:
            writers['market_master'].close()
    for folder, paths in captured.items():
        extension = '_presettlement' if folder == 'candles_presettlement' else ''
        for path in paths:
            try:
                payload = read_json(path)
                source = str(path.relative_to(root))
                for grain, interval in (('hourly', 60), ('daily', 1440)):
                    name = ('api_daily_sessions' if grain == 'daily' else grain) + extension
                    if selected is not None and name not in selected:
                        continue
                    for candle in payload.get(grain) or []:
                        writer(name, CANDLE_COLUMNS).append(candle_row(candle, payload, interval, source, metadata.get(payload.get('market_ticker')), 'presettlement' if extension else 'common'))
                utc_name = 'utc_daily' + extension
                if selected is None or utc_name in selected:
                    for row in utc_daily_rows(payload, source, metadata.get(payload.get('market_ticker')), 'presettlement' if extension else 'common'):
                        writer(utc_name, UTC_DAILY_COLUMNS).append(row)
                for window in payload.get('windows') or []:
                    row = dict(window)
                    row.update({'market_ticker': payload.get('market_ticker'), 'series_ticker': payload.get('series_ticker'), 'dataset': folder, 'raw_artifact': source, 'extra_fields': {k: v for k, v in window.items() if k not in WINDOW_COLUMNS}})
                    writer('window_coverage', WINDOW_COLUMNS).append(row)
            except Exception as exc:
                errors.append({'artifact': str(path.relative_to(root)), 'error': str(exc), 'stage': 'candle_export'})
        for name in ('hourly' + extension, 'api_daily_sessions' + extension, 'utc_daily' + extension):
            if name in writers:
                writers[name].close()
    if selected is None or 'books' in selected:
        for path in captured_books:
            try:
                payload = read_json(path)
                source = str(path.relative_to(root))
                summary, levels = book_rows(payload, source, metadata.get(payload.get('ticker')))
                writer('book_summary', BOOK_SUMMARY_COLUMNS).append(summary)
                for row in levels:
                    writer('book_levels', BOOK_LEVEL_COLUMNS).append(row)
                artifacts.append({'dataset': 'books', 'artifact': source, 'status': payload.get('status'), 'rows': len(levels), 'errors': payload.get('errors')})
            except Exception as exc:
                errors.append({'artifact': str(path.relative_to(root)), 'error': str(exc), 'stage': 'book_export'})
    for row in artifacts:
        writer('source_coverage', ['dataset', 'artifact', 'status', 'rows', 'errors']).append(row)
    for ticker in missing:
        writer('source_coverage', ['dataset', 'artifact', 'status', 'rows', 'errors']).append({'dataset': 'candles', 'artifact': ticker, 'status': 'missing', 'errors': 'No captured candle artifact at export start.'})
    for row in errors:
        writer('export_errors', ['artifact', 'stage', 'error']).append(row)
    files = []
    for item in writers.values():
        item.close()
        files.extend(item.files)
    if errors:
        state = 'PARTIAL EXTRACT — SEE COVERAGE'
    manifest = {'schema_version': 1, 'created_at': utc_now(), 'asof_utc': asof,
                'status': state, 'source_root': str(root), 'output_directory': str(out),
                'public_api_only': True, 'inputs': inputs, 'selected_datasets': sorted(selected) if selected else 'all',
                'files': files, 'rows_by_dataset': {k: v.total for k, v in writers.items()},
                'excel_flags': {k: dict(v.flags) for k, v in writers.items()}, 'errors': errors,
                'source_manifests': docs}
    # The small index contains the scope, limitations and a discoverable workbook inventory.
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    overview = [
        ['Kalshi public API extract', state], ['As of UTC', asof],
        ['Export created UTC', manifest['created_at']], ['Eligible markets', len(seen)],
        ['Markets missing candle artifacts', len(missing)], ['Export/source issues', len(errors)],
        ['Workbook partitions', len(files)], ['Source API', SOURCE_URL],
        ['Calendar windows', f'Eligible since {config.get("eligibility_start_utc", "unavailable")}. Common history: {config.get("hourly_start_utc", "unavailable")} to {asof}. See export_manifest.json for full configuration.'],
        ['Coverage', 'Each output represents captured files at export start. Discovery/category mapping and API errors determine completion.'],
        ['API daily sessions', 'API-defined candle sessions; boundaries may not be UTC midnight. These OHLC, volume and ending OI are raw daily candles.'],
        ['UTC daily aggregation', 'Separate UTC calendar-day aggregation from observed hourly candles. End at 00:00 belongs to previous day. Missing hours and boundary overlaps remain flagged; volume is not prorated.'],
        ['Missing observations', 'Blank is unavailable. No synthetic hourly prices or zero-volume candles are inserted.'],
        ['Notionals', 'Contracts × source face value. Not premium turnover, collateral, directional conviction or unique investors.'],
        ['Large positions', 'Individual trader positions/identity are not available through the public market-data API.'],
        ['Concentration', 'Other participants\' positions and concentration are unavailable through the public market-data API.'],
        ['Large executions', 'Public trade prints can support execution-size research, but exhaustive tick trades are not collected in this bulk candle extract.'],
        ['Order books', 'Historical order-book depth is not supplied by this extract. Any current snapshots are separately timestamped.'],
        ['Book timing', 'Book levels and summaries reflect their own retrieval times after the frozen census cutoff. Displayed orders are not trader positions.'],
        ['Prices', 'USD by field/endpoint schema. Historical plain OHLC is dollars; legacy live plain OHLC is cents.'],
        ['Precision', 'Excel numeric precision is limited to 15 significant digits; exact source strings remain in gzip JSON.'],
        ['Text', 'Formula-like source strings are literal text. XML controls removed and long cells truncated only in Excel, with row flags.'],
        ['Pre-settlement extension', 'Separate presettlement workbooks. Do not pool with common-window rows without aligning windows.'],
        ['Raw data', 'Raw source artifacts remain authoritative for all unflattened fields, exact values and complete rules.'],
        ['Investable universe', 'MSCI ACWI research scope only. No proprietary holdings or benchmark constituent history supplied.'],
        ['Research blueprint', 'research/midterms_equity_blueprint.md'],
    ]
    for row in overview:
        ws.append(row)
    inventory = wb.create_sheet('Files')
    inventory.append(['Dataset', 'Workbook', 'Data rows', 'Sheets', 'Bytes', 'Status'])
    for f in files:
        inventory.append([f['dataset'], f['file'], f['rows'], f['sheets'], f['bytes'], f['status']])
        inventory.cell(inventory.max_row, 2).hyperlink = f['file']
        inventory.cell(inventory.max_row, 2).style = 'Hyperlink'
    dictionary = wb.create_sheet('Dictionary')
    dictionary.append(['Dataset', 'Field', 'Definition'])
    for row in dictionary_rows(columns):
        dictionary.append([row['dataset'], row['field'], row['definition']])
    for sheet in wb:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'B2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='203864')
            cell.font = Font(name='Arial', size=10, color='FFFFFF', bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value, _ = literal(cell.value)
                    cell.data_type = 's'
                cell.font = Font(name='Arial', size=10)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for i in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 28 if i == 1 else 58
    ws.column_dimensions['B'].width = 110
    dictionary.column_dimensions['C'].width = 110
    inventory.column_dimensions['A'].width = 38
    inventory.column_dimensions['B'].width = 58
    for letter in ('C', 'D', 'E'):
        inventory.column_dimensions[letter].width = 17
    for sheet in (ws, inventory):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = '#,##0'
    for sheet in wb:
        for row in sheet.iter_rows():
            lines = max(math.ceil(len(str(cell.value or '')) / max(12, sheet.column_dimensions[cell.column_letter].width - 4)) for cell in row)
            sheet.row_dimensions[row[0].row].height = min(120, max(20, 15 * lines))
    wb.save(out / 'INDEX.xlsx')
    index_check = load_workbook(out / 'INDEX.xlsx', read_only=True, data_only=False)
    if index_check['Summary']['B1'].value != state:
        raise RuntimeError('Saved index state failed verification')
    index_check.close()
    manifest['index'] = 'INDEX.xlsx'
    pending = out / 'export_manifest.json.tmp'
    pending.write_text(json.dumps(manifest, indent=2, ensure_ascii=False, default=str), encoding='utf-8')
    pending.replace(out / 'export_manifest.json')
    print(json_text({'status': state, 'rows': manifest['rows_by_dataset'], 'index': str(out / 'INDEX.xlsx')}), flush=True)
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=Path('data/kalshi_full_20260908'))
    parser.add_argument('--out', type=Path)
    parser.add_argument('--max-sheet-rows', type=int, default=350000)
    parser.add_argument('--max-workbook-rows', type=int, default=500000)
    parser.add_argument('--datasets', nargs='+', choices=['market_master', 'hourly', 'api_daily_sessions', 'utc_daily', 'hourly_presettlement', 'api_daily_sessions_presettlement', 'utc_daily_presettlement', 'books'])
    args = parser.parse_args()
    if not 1 <= args.max_sheet_rows <= 1_048_573:
        parser.error('--max-sheet-rows must be between 1 and 1048573')
    if not 1 <= args.max_workbook_rows <= 2_000_000:
        parser.error('--max-workbook-rows must be between 1 and 2000000')
    if not args.root.exists():
        parser.error(f'Extraction root does not exist: {args.root}')
    out = args.out or args.root / 'excel' / datetime.now(UTC).strftime('%Y%m%dT%H%M%SZ')
    export(args.root, out, args.max_sheet_rows, args.max_workbook_rows, set(args.datasets) if args.datasets else None)


if __name__ == '__main__':
    main()
