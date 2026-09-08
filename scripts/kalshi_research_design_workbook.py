"""Create the proposed public-data research design, without investment scores."""
from __future__ import annotations

import json
import math
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
BLUEPRINT = ROOT / 'research/midterms_equity_blueprint.md'
OUTPUT = ROOT / 'research/Midterms_Research_Design.xlsx'
projects = dict((int(n), title) for n, title in re.findall(r'^## Priority (\d+) — (.+)$', BLUEPRINT.read_text(encoding='utf-8'), re.M))
assert set(projects) == set(range(1, 9)), 'Blueprint must contain the eight agreed priority projects.'

HYPOTHESES = [
    [1, 'Customer protections can make utility load growth more valuable; freezes or cost disallowances can reverse the benefit.', 'Which utilities convert data-center demand into protected 2027–28 cash flows, and who bears the matching cost?', 'State tariffs, contracts and funding', 'Utilities with protected versus speculative large-load demand; costs passed to hyperscalers or tenants.', 'P1, P2'],
    [2, 'Restrictions on new construction can increase rents on existing connected capacity while impairing greenfield projects.', 'Where do existing connection rights and near-term contracted capacity outweigh pipeline size?', 'Local siting and grid access', 'Connected assets versus speculative pipeline; equipment orders that survive delay or relocation.', 'P3, P4'],
    [3, 'Price caps, subsidies and supply expansion transfer costs differently; household cash relief need not match inflation relief.', 'Who benefits from more residual household spending, and which companies or taxpayers fund it?', 'Regional household budgets and policy incidence', 'Within-industry consumer, lender, builder and landlord comparisons by customer geography.', 'P5, P6'],
    [4, 'Expected Fed cuts can coexist with higher term premiums, mortgage rates or credit spreads.', 'Which companies need long-term funding relief rather than lower short rates?', 'Expected rates, term premium and credit', 'Debt-maturity and deposit-sensitivity comparisons within utilities, banks, builders and real estate.', 'P7, P8'],
    [5, 'Falling gasoline due to restored supply has different equity effects from falling gasoline due to demand destruction.', 'Which producer/consumer mix is resilient to persistent supply disruption and weak demand?', 'Energy supply, voter budgets and corporate margins', 'Fuel producers, refiners, transport and consumer firms, accounting for hedges and pass-through.', 'P9, P10, P11'],
    [6, 'Software harmonization can reduce compliance fragmentation while state/local infrastructure constraints tighten.', 'Where can AI adoption scale economically despite use-case liability and infrastructure restrictions?', 'Federal proposals, state duties and adoption', 'Model developers, deployers and incumbents with different regulated-use exposure.', 'P12, P13'],
    [7, 'Executable supply reform and demand subsidies can have opposing effects on builders, land values and incumbent rents.', 'Which local policy mix favors construction volumes, unit margins or incumbent scarcity rents?', 'Permitting, subsidies and housing supply', 'Metro-matched builders/materials versus landlords, with mortgage and migration controls.', 'P14'],
    [8, 'State Medicaid execution, risk mix and reimbursement lags can matter more to earnings than national party labels.', 'Which managed-care or hospital exposures are most sensitive to 2027–28 implementation?', 'Federal requirements and state administration', 'Membership, revenue-per-member, medical-cost and hospital payer-mix sensitivities.', 'P15'],
]

TRANSMISSION = [
    [1, 'State utility commission; legislature; utility and large-load customer.', 'Governor/legislature may influence appointments, statute or political pressure; verify staggered commission terms.', 'Rate-case/tariff order; customer contract; cost-recovery approval.', 'Contract enforceability, collateral, actual connection and allowed cost recovery.', 'Minimum revenue, rate base, capital needs, allowed return and counterparty losses.', 'Map decision dates and commissioning to 2027–28; existing contract protections may matter before new laws.', 'AEP/PUCO precedent is an evidence anchor, not proof about every jurisdiction.'],
    [2, 'County/municipality; state bodies; grid operator; federal regulator where applicable.', 'Relevant governor, legislature, local or referendum outcomes; record races absent from Kalshi.', 'Zoning, permit, connection, moratorium or grandfathering decision.', 'Grid and construction bottlenecks; substitution; lease reset and cancellation terms.', 'Occupied rent, energization delay, backlog revenue, working capital and stranded capital.', 'Use project critical paths; explicitly test assumed 6-, 12- and 24-month delays.', 'Queue applications are not equivalent to economically committed demand.'],
    [3, 'Federal/state legislature; regulator; local authority according to the proposal.', 'Identify which elections can change each subsidy, cap, tax or supply measure.', 'Verify funding, votes, eligibility, fiscal offsets and legal scope.', 'Administrative delivery, price pass-through and physical supply response.', 'Household residual spending; company volume, margin, credit loss and payer burden.', 'Immediate cash transfers can precede slower supply effects; underwriting horizon is 12–24 months.', 'CPI/PCE inflation, rent and household mortgage cash payments are different measures.'],
    [4, 'FOMC; fiscal authorities; debt issuers; financial markets.', 'Election outcomes inform fiscal/policy branches; they do not directly set the Fed rate.', 'Separate fiscal enactment from monetary decisions and term-premium repricing.', 'Debt repricing, refinancing, deposit adjustment, mortgage transmission and demand.', 'Interest expense, net interest income, credit losses, customer affordability and valuation.', 'Align each Fed meeting and debt maturity; do not extrapolate a near-term cut into all 2028 funding costs.', 'Marginal rate/election contracts do not identify conditional paths or joint probabilities.'],
    [5, 'Energy producers, logistics and external governments; U.S. policy authority varies by action.', 'Gasoline can affect voting while geopolitics simultaneously affects both elections and profits.', 'Identify actual energy, tax, reserve or trade actions; do not substitute campaign language.', 'Shipping/production restoration; inventories; refinery response; hedging and pump-price pass-through.', 'Real household budgets; commodity realizations; refining margin; unhedged fuel costs and volumes.', 'Separate election-day prices, shock persistence, hedge maturity and 2027–28 earnings.', 'Supply restoration and recessionary demand weakness must be distinct scenarios.'],
    [6, 'Congress; federal agencies/courts; state legislatures, attorneys general and sector regulators.', 'Chamber/state outcomes alter feasibility; a proposed federal framework is not enacted preemption.', 'Audit enacted duties, exemptions, liability, procurement and any valid federal preemption.', 'Compliance build, legal challenges, customer procurement and application deployment.', 'Adoption/revenue timing, compliance cost, margin and liability exposure.', 'Map effective dates and customer sales cycles to the investment horizon.', 'Model developers, deployers and infrastructure owners face different obligations.'],
    [7, 'Local planning bodies; state legislature/governor; funding and housing agencies.', 'Separate state election leverage from local implementation and referendum authority.', 'Supply/permitting reform; subsidies; rent rules; tax and infrastructure funding.', 'Permits, starts, completions, land exercise, infrastructure and finance availability.', 'Unit volumes/margins, land option value, materials demand, rents and concessions.', 'Existing owned land may reprice before completions; supply takes time.', 'A policy precedent is not evidence of the probability of a new law passing.'],
    [8, 'Federal statute/CMS; state Medicaid agency; managed-care contracting and providers.', 'State elections may change execution choices, not erase federal statutory requirements.', 'Verify federal requirements, state discretion, waivers, rates and administrative decisions.', 'Enrollment procedures, outreach, eligibility systems, appeals and reimbursement lags.', 'Members × revenue per member × medical-cost economics; provider payer mix and uncompensated care.', 'CMS identifies a January 2027 implementation channel; verify state-specific transitions.', 'Lower membership does not automatically imply higher insurer margins.'],
]

DATA = [
    ['Universe', 'Kalshi public API', 'Series, event and individual market identifiers; rules; strikes; lifecycle; timestamps.', 'Contract plus category crosswalk', 'Every eligible contract, including thin/zero volume; avoid double-counting overlapping category paths.', 'Exact historical website category assignment may not be publicly reconstructible.'],
    ['Electoral outcomes', 'Kalshi public API', 'Federal/state race and control prices with rules and outcome orientation.', 'Market and synchronized observation', 'Classify exclusive buckets, cumulative thresholds, conjunctions and duplicate propositions.', 'Marginals alone cannot recover election-policy-macro conditionals.'],
    ['Hourly history', 'Kalshi public API', 'Raw Yes bid/ask/price OHLC; volume; ending open interest.', 'Hourly API candle', 'USD units by endpoint schema; fractional contracts; preserve sparse rows and nulls.', 'OHLC is not intrahour trade chronology or historical order-book depth.'],
    ['API daily sessions', 'Kalshi public API', 'Raw daily OHLC, volume and ending OI.', 'API session start/end', 'Preserve actual boundaries; do not label an API session as UTC midnight-to-midnight.', 'Partial requested-window boundary candles are not prorated.'],
    ['UTC daily coverage', 'Derived from public Kalshi hourly data', 'Observed volume sum; latest OI and timestamp; observed/expected hours; boundary flags.', 'UTC date, using candle end minus one second', 'OI is a stock and is never summed. Incomplete observations remain identified.', 'Missing hours cannot be silently treated as zero trading.'],
    ['Current books', 'Kalshi public API', 'Yes/No bid ladders, displayed sizes and retrieval times.', 'Retrieval snapshot after census cutoff', 'Derive opposite ask from face minus bid; retain original side and quantity.', 'Displayed orders are not participant positions; historical depth unavailable.'],
    ['Positions and concentration', 'Unavailable through public market-data API', 'Other traders’ identities, complete positions and concentration.', 'Not available', 'Record unavailable, not zero. Do not infer identity or motivation from size.', 'Public executions can be studied separately; exhaustive ticks are not in the bulk candle extract.'],
    ['Policy authority and status', 'Public primary sources', 'Enacted text, proposals, official orders, commission terms, effective dates and court status.', 'Jurisdiction and decision', 'Maintain version and retrieval dates; distinguish proposal, enactment and implementation.', 'Requires manual legal-scope validation; party labels are insufficient.'],
    ['Utility customer protections', 'Public tariffs, commission dockets and company filings', 'Minimum bills, collateral, exit payments, connection costs, rate base and funding.', 'Service territory, tariff and contract', 'Compare source provisions and asset-specific cash-flow exposure.', 'Some contract details may be confidential; missing terms stay missing.'],
    ['Data-center projects', 'Public permits, grid/utility documents and filings', 'Location, energization rights, grandfathering, leases, orders and project milestones.', 'Project and issuer', 'Separate applications, committed projects and commissioned assets.', 'Public queue data can overstate deliverable or economically committed demand.'],
    ['Household affordability', 'BLS, BEA, EIA and other public agencies', 'Expenditure weights; rent; energy; employment; income; housing-finance cash costs.', 'Publicly supported household/geographic cohort', 'Keep CPI/PCE and cash budgets separate; store release and reference dates.', 'Geographic granularity and consistent cohort histories can be limited.'],
    ['Fed and financing', 'Federal Reserve, Treasury and New York Fed public data', 'Policy decisions, yield curve, term-premium estimates and public financing indicators.', 'Meeting, maturity and publication date', 'Separate expected short rates, term premium, credit and mortgage costs.', 'Term-premium decomposition is model-based; public high-frequency coverage may be limited.'],
    ['Energy regimes', 'EIA and public company/market disclosures', 'Production, inventories, demand, refinery data, shipping evidence, hedging and fuel costs.', 'Release, region and issuer', 'Distinguish supply restoration, supply disruption and demand destruction.', 'Public futures/history access may limit synchronized event studies.'],
    ['AI use-case exposure', 'Enacted laws, official guidance and public filings', 'Developer/deployer duties, regulated applications, exemptions, procurement and compliance cost.', 'Jurisdiction, use case and issuer', 'Audit legal reach and deployment timing before assigning earnings exposure.', 'A mention contract measures expected wording, not revenue or sentiment.'],
    ['Housing execution', 'Public planning/permit data, agencies and filings', 'Land/options, permits, starts, concessions, rents, insurance and migration.', 'Metro, asset and issuer', 'Separate supply reforms, subsidies and rent regulation.', 'Comparable asset-level disclosures may be unavailable.'],
    ['Medicaid execution', 'CMS, state agencies and public insurer/provider filings', 'Membership, rates, medical costs, risk mix, provider payer mix and rollout milestones.', 'State, plan/provider and reporting quarter', 'Separate federal statute, state discretion and administrative execution.', 'State/issuer disclosure may not support precise local attribution.'],
    ['Company economics and valuation', 'Public annual/interim filings, guidance and investor materials', 'Revenue, margins, tax, debt, capex, working capital, share count and dated market prices.', 'Issuer, currency and scenario', 'Build 2027–28 cash flows and valuation assumptions separately.', 'No proprietary consensus. Public guidance and labeled assumptions replace consensus claims.'],
    ['Global security mapping', 'Public issuer/exchange disclosures; user-supplied data if later provided', 'Share class, listing, domicile, currency, production, customers and economic FX exposure.', 'Security and issuer', 'Separate country allocation, currency effects and within-industry security selection.', 'No licensed ACWI constituent history, official weights or actual holdings supplied.'],
    ['Mentions and attention', 'Kalshi plus public transcripts and policy documents', 'Expected words, actual transcripts and subsequent verifiable actions.', 'Stable speaker/event cohort', 'Preregister vocabulary and normalize for event type and listing selection.', 'Mention probability is not voter support, policy commitment or an earnings forecast.'],
    ['Implementation costs', 'Public quotes and disclosures where available', 'Liquidity, spreads, options terms and borrow/carry inputs for proposed expressions.', 'Security, trade and expiry', 'Evaluate cost, catalyst alignment and residual exposures before implementation.', 'If public borrow/options inputs are missing, feasibility remains unverified.'],
]

FALSIFICATIONS = [
    [1, 'Contract advantage', 'Read tariff/contract cancellation, collateral and cost-recovery provisions.', 'Counterparties can exit cheaply, protections are unenforceable or costs are disallowed.', 'Do not rank by announced megawatts alone.'],
    [1, 'Incremental value', 'Compare 2027–28 protected cash flows and financing with a dated public valuation baseline.', 'Protection is outside the horizon or fully reflected in value.', 'A superior business exposure is not automatically a mispriced stock.'],
    [2, 'Scarcity economics', 'Compare connected and greenfield assets within the same grid region.', 'Grandfathering protects projects, substitutes appear or fixed leases prevent rent capture.', 'Control for actual demand and relocation.'],
    [2, 'Delay sensitivity', 'Translate project delays into revenue, working capital and funding under explicit assumptions.', 'Orders remain noncancelable or delays fall outside the relevant cash-flow horizon.', 'Do not estimate a delay effect from a political price move alone.'],
    [3, 'Household signal', 'Test whether regional cash-budget stress adds value beyond income, jobs and national inflation.', 'No stable incremental signal or exposures cannot be mapped.', 'Use longer public histories for household elasticities.'],
    [3, 'Cost incidence', 'Trace household relief to corporate margin, fiscal funding, taxes and supply.', 'Relief is offset by other bills, taxes, reduced quality or impaired supply.', 'Show the payer and beneficiary in the same scenario.'],
    [4, 'Financing channel', 'Inspect debt maturities, hedges, deposit pricing and customer borrowing.', 'Fixed funding/hedges make the proposed rate effect immaterial.', 'Separate short-rate, duration and credit effects.'],
    [4, 'Political attribution', 'Compare political-news windows with CPI, payroll and central-bank windows.', 'The association disappears with common-news controls or term-premium specification changes.', 'Two months cannot identify a long-horizon causal return model.'],
    [5, 'Energy regime', 'Cross-check inventory, activity, production and shipping evidence.', 'Observed supply or demand contradicts the proposed regime.', 'Geopolitics can be a common cause of election and earnings changes.'],
    [5, 'Firm sensitivity', 'Audit fuel hedges, refinery configuration, pricing and customer volumes.', 'Hedging/pass-through eliminates the impact or producer protection fails under rent limits.', 'Do not treat all energy or transport stocks as equivalent.'],
    [6, 'Legal reach', 'Verify enacted duties, valid preemption, exemptions and court status.', 'Applicable law supersedes the modeled duty, exemptions remove exposure or timing is too late.', 'A federal proposal is not enacted blanket preemption.'],
    [6, 'Commercial effect', 'Match regulated use-case exposure to adoption timing and compliance costs.', 'Disclosures cannot support mapping or ordinary IT demand/product quality explains results.', 'Separate developer, deployer and infrastructure economics.'],
    [7, 'Policy classification', 'Separate supply expansion, demand subsidy, rent rules and infrastructure.', 'Assumed supply reform cannot deliver permits/starts within the horizon.', 'Test mortgages, migration, insurance, labor and local pre-trends.'],
    [7, 'Asset economics', 'Map actual land/options and rental assets to affected metros.', 'Land exposure, concessions or rental outcomes contradict the proposed trade.', 'Lower house prices need not improve every builder’s margin.'],
    [8, 'Execution and discretion', 'Verify state readiness, statutory constraints, rollout dates and comparability.', 'States lack modeled discretion or comparison groups fail pre-trends.', 'Do not infer implementation solely from party control.'],
    [8, 'Membership and margin', 'Model risk mix, medical cost, reimbursement and provider payer mix jointly.', 'Enrollment decline raises medical-cost pressure or hospital losses beyond assumed offsets.', 'Lower membership is not a sufficient insurer-profit thesis.'],
    [None, 'Market quality', 'Check synchronized quote/trade evidence, spreads, sparse hours and lifecycle.', 'Price moves are stale, nonexecutable, definition changes or a coverage artifact.', 'Keep price source, liquidity and timestamp flags.'],
    [None, 'Probability identification', 'Audit outcome partitions and any conditional contract definitions.', 'The desired joint/conditional probability is unidentified by observed marginals.', 'Use explicit bounds or documented scenario assumptions; do not multiply marginals automatically.'],
    [None, 'Forecast validation', 'Use comparable resolved events, event clustering and prospective/out-of-sample logs.', 'Apparent performance is driven by repeated snapshots, easy daily outcomes or revised data.', 'Separate election, policy and investment forecast accuracy.'],
    [None, 'Decision gate', 'Require exposure, timing, value gap, downside and implementability.', 'A significant coefficient or odds move has no material company-value consequence.', 'Publish negative findings and missing inputs.'],
]

SOURCES = [
    ['P1', 1, 'PUCO data-center tariff announcement', 'https://content.govdelivery.com/accounts/OHPUC/bulletins/3e8bb79', 'Evidence anchor: July 2025 Ohio customer-protection tariff decision.'],
    ['P2', 1, 'AEP Ohio data-center tariff', 'https://www.aepohio.com/company/about/rates/data-center-tariff/', 'Tariff provisions and implementation reference.'],
    ['P3', 2, 'Virginia JLARC data-center recommendations', 'https://jlarc.virginia.gov/pdfs/other/Unimp-Leg-Recs/2025DataCenters.pdf', 'Local/state siting and infrastructure channels.'],
    ['P4', 2, 'FERC large-load integration action', 'https://www.ferc.gov/news-events/news/ferc-launches-aggressive-targeted-action-speed-large-load-integration', 'Separate federal/grid authority; June 2026 evidence anchor.'],
    ['P5', 3, 'BLS rent and owners’ equivalent rent methodology', 'https://www.bls.gov/cpi/factsheets/owners-equivalent-rent-and-rent.htm', 'Housing services differ from home-financing cash payments.'],
    ['P6', 3, 'BEA CPI–PCE differences', 'https://www.bea.gov/index.php/help/faq/555', 'Price-index formula, weights and scope.'],
    ['P7', 4, 'New York Fed Treasury term-premium data', 'https://www.newyorkfed.org/research/data_indicators/term-premia-tabs', 'Model-based expected-short-rate and term-premium decomposition.'],
    ['P8', 4, 'Federal Reserve legal framework', 'https://www.federalreserve.gov/faqs/basic-framework-determines-conduct-monetary-policy.htm', 'Statutory goals and operational independence.'],
    ['P9', 5, 'June 2026 FOMC minutes', 'https://www.federalreserve.gov/monetarypolicy/fomcminutes20260617.htm', 'Historical energy/AI inflation transmission evidence, not a current military-status assertion.'],
    ['P10', 5, 'EIA gasoline-price components', 'https://www.eia.gov/energyexplained/gasoline/factors-affecting-gasoline-prices.php', 'Crude, refining, distribution and tax channels.'],
    ['P11', 5, 'EIA Short-Term Energy Outlook', 'https://www.eia.gov/outlooks/steo/report/', 'Use a dated vintage and distinguish forecast assumptions from realized outcomes.'],
    ['P12', 6, 'White House March 2026 AI legislative framework', 'https://www.whitehouse.gov/releases/2026/03/president-donald-j-trump-unveils-national-ai-legislative-framework/', 'A framework calling for legislation; not proof of enacted blanket preemption.'],
    ['P13', 6, 'Colorado SB26-189', 'https://www.leg.colorado.gov/bills/SB26-189', 'State law and implementation framework.'],
    ['P14', 7, 'California 2025 housing reform enactment', 'https://www.gov.ca.gov/2025/06/30/governor-newsom-signs-into-law-groundbreaking-reforms-to-build-more-housing-affordability/', 'Policy precedent for supply reform; not a prediction of 2026 enactment.'],
    ['P15', 8, 'CMS community engagement implementation', 'https://www.medicaid.gov/resources-for-states/working-families-tax-cut-legislation/community-engagement', 'Federal framework and January 2027 implementation channel.'],
    ['D1', None, 'Kalshi public API documentation', 'https://docs.kalshi.com/', 'Endpoint definitions, authentication/access, units and response fields.'],
    ['D2', None, 'Kalshi fixed-point migration', 'https://docs.kalshi.com/getting_started/fixed_point_migration', 'Fractional contracts and price-field units.'],
    ['D3', None, 'Kalshi order-book response conventions', 'https://docs.kalshi.com/getting_started/orderbook_responses', 'Public Yes/No bid ladders and opposite-ask relationship.'],
    ['R1', None, 'Primary research blueprint', 'midterms_equity_blueprint.md', 'Authoritative eight-project priority and full methodology.'],
    ['R2', None, 'Macro research designs', 'macro_research_ideas.md', 'Macro hypotheses, identification and ACWI attribution.'],
    ['R3', None, 'State-policy research designs', 'policy_equity_ideas.md', 'Authority, state/local implementation and equity transmission.'],
    ['R4', None, 'Kalshi API feasibility', 'kalshi_api_feasibility.md', 'Public-data availability and census limitations.'],
    ['R5', None, 'Company candidate map', 'company_candidate_map.md', 'Illustrative candidates for diligence; does not imply scored holdings or recommendations.'],
]


def add_table(wb, name, title, subtitle, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws['A1'] = title
    ws['A1'].font = Font(name='Arial', size=14, bold=True, color='203864')
    ws['A2'] = subtitle
    ws['A2'].font = Font(name='Arial', size=10, italic=True, color='555555')
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for i, header in enumerate(headers, 1):
        ws.cell(4, i, header)
    for row in rows:
        ws.append(row)
    table = Table(displayName=name.replace('_', '') + 'Table', ref=f'A4:{get_column_letter(len(headers))}{4 + len(rows)}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    ws.freeze_panes = 'C5' if len(headers) > 3 else 'B5'
    ws.auto_filter.ref = table.ref
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = 's'
            elif isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0'
        lines = max(math.ceil(len(str(c.value or '')) / max(12, widths[c.column - 1] - 4)) for c in row)
        ws.row_dimensions[row[0].row].height = min(180, max(30, 15 * lines + 8))
    for cell in ws[4]:
        cell.fill = PatternFill('solid', fgColor='203864')
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)
    ws.row_dimensions[1].height = 23
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 7
    ws.print_title_rows = '1:4'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)
    subtitle = 'Proposed research, not tested findings. Public data. Global ACWI framework. Horizon: 12–24 months.'
    rows = [[p, projects[p], hypothesis, question, channel, output, evidence, 'Proposed; empirical work pending'] for p, hypothesis, question, channel, output, evidence in HYPOTHESES]
    add_table(wb, 'Priority_Hypotheses', 'Eight priority research projects', subtitle,
              ['Priority', 'Project', 'Falsifiable hypothesis', 'Investment question', 'Transmission channel', 'Proposed portfolio output', 'Source IDs', 'Evidence status'],
              rows, [10, 42, 61, 55, 32, 62, 16, 29])
    add_table(wb, 'Transmission_Map', 'From election outcomes to company cash flows', 'Each link requires evidence. Election probabilities do not supply enactment or implementation probabilities.',
              ['Priority', 'Project', 'Decision-making authority', 'Election link', 'Enactment or decision', 'Implementation dependency', 'Cash-flow channel', 'Timing to test', 'Key distinction'],
              [[r[0], projects[r[0]], *r[1:]] for r in TRANSMISSION], [10, 36, 40, 49, 48, 52, 47, 52, 50])
    add_table(wb, 'Data_Requirements', 'Required observations and data limits', 'Preserve source definitions, publication dates, units, missing values and collection coverage.',
              ['Workstream', 'Source/access', 'Required fields', 'Grain', 'Use and handling', 'Known limitation'], DATA, [29, 47, 67, 38, 68, 67])
    add_table(wb, 'Falsification_Tests', 'Tests that can reject an investment thesis', 'Negative findings and uncertain mappings remain part of the research record.',
              ['Priority', 'Project', 'Test', 'Proposed evidence', 'Rejection condition', 'Interpretation discipline'],
              [[r[0], projects.get(r[0], 'All projects'), *r[1:]] for r in FALSIFICATIONS], [10, 38, 30, 66, 65, 62])
    expressions = []
    for p, hypothesis, question, channel, proposed, evidence in HYPOTHESES:
        expressions += [
            [p, projects[p], 'Long-only', proposed, 'Fund overweights through underweights in comparable benchmark exposures; no assumed weights.', '2027–28 cash-flow and value gap; benchmark-relative scenario loss.', 'Check sector, country, currency, size, valuation and market-beta exposures; actual holdings/official weights unavailable.'],
            [p, projects[p], '130/30', 'Within-industry long/short expression of the same economic mechanism, if the public evidence supports it.', 'Short side must have a weaker valuation-adjusted payoff, not merely a different political label.', 'Incremental expected value after financing, borrow, dividends, spreads and liquidity.', 'Borrow availability/cost and squeeze exposure must be checked before implementation; no fabricated position sizing.'],
            [p, projects[p], 'Options or hedge', 'Compare suitable stock, sector, rate, commodity or currency hedges for the modeled loss channel.', 'Choose expiry and instrument to cover the actual catalyst and cash-flow timing.', 'Scenario-loss reduction per unit of premium/carry and active risk; evaluate basis and path risk.', 'An election-expiry contract does not automatically hedge a policy cash-flow effect realized in 2028; public quotes may be incomplete.'],
        ]
    add_table(wb, 'Portfolio_Expressions', 'Translate findings into benchmark-relative decisions', 'Generic ACWI expressions. No holdings, company scores, forecast probabilities, position sizes or trade recommendations are fabricated.',
              ['Priority', 'Project', 'Implementation', 'Economic expression', 'Funding or construction', 'Decision metric', 'Required attribution and feasibility'],
              expressions, [10, 38, 22, 67, 65, 61, 76])
    sources = add_table(wb, 'Sources', 'Evidence anchors and research files', 'Primary sources substantiate definitions and precedents. They do not establish the proposed investment conclusions.',
                        ['Source ID', 'Priority', 'Source', 'URL or relative research file', 'Use and caveat'], SOURCES, [13, 10, 49, 75, 80])
    for row in range(5, sources.max_row + 1):
        link = sources.cell(row, 4)
        if str(link.value).startswith('http') or (OUTPUT.parent / str(link.value)).exists():
            link.hyperlink = link.value
            link.font = Font(name='Arial', size=10, color='0563C1', underline='single')
    scope = [
        ['Research status', 'Proposed hypotheses and tests. No completed empirical stock-selection findings in this workbook.'],
        ['Prepared', datetime(2026, 9, 8)],
        ['Priority order', 'The eight projects retain the order in the primary research blueprint. State policy, AI/data centers and affordability receive the deepest work.'],
        ['Mandate', 'Global equity portfolio framework benchmarked to MSCI ACWI, approximately 3,000 securities as described by the user.'],
        ['Investment horizon', '12–24 months; focus on 2027–28 company earnings, cash flows and valuation.'],
        ['Implementation', 'Long-only and 130/30. Shorts and options are permitted in the research design, subject to implementability.'],
        ['Data access', 'Public API and public external evidence only. No proprietary holdings, forecasts, consensus or institutional risk model supplied.'],
        ['Reporting currency', 'USD for illustrations, as specified in the blueprint; separate local-currency operating exposure and portfolio translation.'],
        ['Kalshi scope', 'All eligible individual contracts in the 33 requested overlapping categories, including thin and zero-volume contracts.'],
        ['Eligibility', 'Active/unresolved contracts plus contracts settled within the preceding three calendar months, using frozen run-config UTC bounds.'],
        ['Common history', 'Hourly history for the preceding two calendar months, or inception if later. Keep market lifecycle and sparse observations explicit.'],
        ['Resolved extension', 'Separate final-two-months-before-settlement dataset for eligible recently settled contracts otherwise lacking common-window history.'],
        ['Daily distinction', 'Raw API daily sessions and separately derived UTC daily observations. Never sum OI; annotate missing hours and full boundary candles.'],
        ['Snapshot timing', 'Current metadata and books reflect documented retrieval times after the frozen census cutoff. They are not historical quote snapshots at that cutoff.'],
        ['Category membership', 'Preserve API category, queries, tags and website-path evidence separately. Exact historical website membership may remain approximate.'],
        ['Private positions', 'Other participants’ identities, positions and concentration are unavailable through the public market-data API, not zero.'],
        ['Trades and books', 'Bulk extraction supplies candles and current books where collected. Exhaustive trade ticks and historical full depth are not promised.'],
        ['Probability limits', 'Marginals do not determine joint or conditional probabilities. Policy-tree transitions and dependence need explicit evidence or labeled assumptions.'],
        ['Validation horizon', 'Two months of market history cannot validate a 12–24-month causal equity-return model. Use earlier public episodes and a prospective record.'],
        ['Public valuation baseline', 'Use dated public prices, filings and guidance with explicit modeling assumptions. Do not claim access to institutional consensus.'],
        ['Portfolio attribution', 'Separate within-industry selection, industry/country allocation, currency and broad style/market sensitivity. Economic location differs from domicile.'],
        ['Missing universe inputs', 'No actual portfolio holdings, licensed constituent history or official historical ACWI weights supplied. A generic framework is not a live portfolio diagnosis.'],
        ['Decision gate', 'Material economic exposure, credible timing, a value gap, acceptable downside and implementable liquidity/costs must all be supported.'],
        ['Primary document', 'midterms_equity_blueprint.md'],
        ['Candidate diligence map', 'company_candidate_map.md; illustrative diligence candidates only, subject to the file’s coverage/status.'],
        ['Dataset completion', 'This planning workbook makes no claim that extraction is complete. Read the dataset INDEX.xlsx and export_manifest.json for current coverage.'],
    ]
    scope_ws = add_table(wb, 'Scope', 'Mandate, definitions and boundaries', 'Keep this scope distinct from any subsequent empirical result or data-collection completion claim.',
                         ['Item', 'Agreed scope or limit'], scope, [34, 132])
    scope_ws['B6'].number_format = 'yyyy-mm-dd'
    for sheet in wb:
        sheet.sheet_properties.tabColor = '203864' if sheet.title in ('Priority_Hypotheses', 'Portfolio_Expressions') else '708090'
    wb.properties.title = 'Midterms Research Design'
    wb.properties.subject = 'Public-data research hypotheses for a global fundamental equity portfolio'
    wb.properties.creator = 'Research planning'
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    checked = load_workbook(OUTPUT, data_only=False)
    assert len(checked.sheetnames) == 7
    assert checked['Priority_Hypotheses'].max_row - 4 == 8
    assert checked['Transmission_Map'].max_row - 4 == 8
    assert checked['Portfolio_Expressions'].max_row - 4 == 24
    assert all(len(ws.tables) == 1 and ws.freeze_panes for ws in checked)
    assert all(cell.data_type != 'f' for ws in checked for row in ws for cell in row)
    report = {'file': str(OUTPUT), 'status': 'proposed_research_not_findings',
              'sheets': {ws.title: ws.max_row - 4 for ws in checked},
              'validation': 'Saved/reopened; eight priorities preserved; seven filterable tables; literal text; no fabricated numeric forecasts.'}
    checked.close()
    (OUTPUT.parent / 'Midterms_Research_Design_validation.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps(report))


if __name__ == '__main__':
    main()
