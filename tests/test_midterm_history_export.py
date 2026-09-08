"""Daily summaries preserve flows, last observed stocks, UTC boundaries and gaps."""

from datetime import datetime

from scripts.export_midterm_history import DAY_FIELDS, Shards, day_ts, summarize_day, utc_cell


def test_midnight_end_belongs_to_prior_utc_day_and_excel_datetime_is_naive():
    assert day_ts(86400) == 0
    assert day_ts(86401) == 86400
    assert utc_cell(86400) == datetime(1970, 1, 2)
    assert utc_cell(86400).tzinfo is None


def test_excel_shards_preserve_numeric_datetime_and_inert_formula_text(tmp_path):
    from openpyxl import load_workbook

    shards = Shards(tmp_path, "HOURLY", ["hour_end_UTC", "value", "title"], 2, notes=[])
    for n in range(3):
        shards.append([utc_cell(3600 * n), 1.25, '=HYPERLINK("https://example.org")'])
    shards.close()
    assert [record["rows"] for record in shards.files] == [2, 1]
    book = load_workbook(tmp_path / "HOURLY_001.xlsx", read_only=True, data_only=False)
    row = list(book["HOURLY"].iter_rows(min_row=2, max_row=2))[0]
    assert row[0].value == datetime(1970, 1, 1)
    assert row[1].value == 1.25
    assert row[1].data_type == "n"
    assert row[2].data_type == "s"
    book.close()


def test_volume_sums_but_oi_uses_last_nonmissing_and_missing_stays_missing():
    market = {
        "ticker": "T",
        "series_ticker": "S",
        "event_ticker": "E",
        "channel": "policy",
        "issue_ids": "energy_iran",
        "face_usd": 2,
        "close_time": "1970-01-01T03:00:00Z",
    }
    rows = [
        {
            "end_ts": 3600,
            "volume_contracts": 4,
            "open_interest_contracts": 10,
            "usable_quote": True,
            "mid_probability": 0.4,
            "boundary_partial": True,
        },
        {
            "end_ts": 7200,
            "volume_contracts": None,
            "open_interest_contracts": 15,
            "usable_quote": False,
            "mid_probability": None,
            "boundary_partial": False,
        },
        {
            "end_ts": 10800,
            "volume_contracts": 6,
            "open_interest_contracts": None,
            "usable_quote": True,
            "mid_probability": 0.6,
            "boundary_partial": False,
        },
    ]
    result = dict(zip(DAY_FIELDS, summarize_day(market, rows, 86400, 900, 0, 86400)))
    assert result["daily_volume_contracts"] == 10
    assert result["volume_observed_intervals"] == 2
    assert result["OI_last_observed_contracts"] == 15
    assert result["OI_face_exposure_USD"] == 30
    assert result["last_usable_midpoint_probability"] == 0.6
    assert result["incomplete_UTC_day"] is True
    assert result["missing_or_outside_nominal_24_intervals"] == 21
    assert result["report_eligible_intervals"] == 2
    assert result["report_daily_volume_contracts"] == 4
    assert result["report_daily_last_OI_contracts"] == 15
    missing = [{**rows[0], "volume_contracts": None, "open_interest_contracts": None}]
    result = dict(zip(DAY_FIELDS, summarize_day(market, missing, 86400, 900, 0, 86400)))
    assert result["daily_volume_contracts"] is None
    assert result["OI_last_observed_contracts"] is None
    assert result["OI_face_exposure_USD"] is None
